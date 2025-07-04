import os
import csv
import json
import re
import unicodedata
import string
import numpy as np
from typing import Optional
import PyPDF2
import pdfplumber
from docx import Document
import openpyxl
from sentence_transformers import SentenceTransformer
from collections import defaultdict
from sentence_transformers import util
import logging_function
import logging
import time
from environment_vars import EMBEDDING_MODEL, EMBEDDINGS_DIR, CHUNKS_DIR, SOURCE_DIR

logger = logging.getLogger(__name__)

patterns = [
    r'page\s*\d+',
    r'daggerheart srd',
    r'https?://\S+',
    r'\b[\w\.-]+@[\w\.-]+\.\w+\b',
    r'\[\d+\]',
    r'\bsection\s*\d+(\.\d+)*\b',
    r'[-_]{2,}',
    r'\*{2,}',
    r'^\d+\.\s+',
    r'\xa0',
]

def read_headings_csv(csv_path):
    headings = []
    with open(csv_path, newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            # Use the lowest-level heading available
            for level in ['subsubsubsubsection', 'subsubsubsection', 'subsubsection', 'subsection', 'section']:
                if row.get(level) and row[level].strip():
                    row['lowest_heading'] = row[level].strip()
                    break
            else:
                row['lowest_heading'] = None
            headings.append(row)
    return headings


# ========== PREPROCESSING ==========
def preprocess_and_get_page_map(text):
    lines = text.split('\n')
    clean_lines = []
    page_map = {}
    current_page = 1
    for i, line in enumerate(lines):
        # Match footer like "Daggerheart SRD,### Page 3" or "Daggerheart SRD, Page 3"
        m = re.match(r".*Daggerheart SRD.*Page\s*(\d+)", line)
        if m:
            current_page = int(m.group(1))
            page_map[len(clean_lines)] = current_page  # Map next line to this page
            continue  # Skip this line
        # Also remove lines that are just "Daggerheart SRD"
        if line.strip().startswith("Daggerheart SRD"):
            continue
        clean_lines.append(line)
    return "\n".join(clean_lines), page_map

def preprocess_text(text, remove_patterns=None):
    logger.debug("preprocess_text called with remove_patterns:", remove_patterns)
    if remove_patterns is None:
        remove_patterns = [
            r'page\s*\d+',
            r'daggerheart srd',
            r'https?://\S+',
            r'\b[\w\.-]+@[\w\.-]+\.\w+\b',
            r'\[\d+\]',
            r'\bsection\s*\d+(\.\d+)*\b',
            r'[-_]{2,}',
            r'\*{2,}',
            r'^\d+\.\s+',
            r'\xa0',
        ]
    text = text.lower()
    text = ''.join(filter(lambda x: x in string.printable, text))
    text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode('utf-8')
    if remove_patterns:
        for pat in remove_patterns:
            text = re.sub(pat, '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    text = text.replace('“', '"').replace('”', '"').replace('’', "'").replace('–', '-').replace ('\n', ' ')
    return text

# ========== TXT EXTRACTION ==========
def extract_text_from_txt(txt_path: str) -> str:
    with open(txt_path, 'r', encoding='utf-8') as file:
        return file.read()

# ========== PDF PAGE EXTRACTION ==========

def is_subsub_heading(line):
    # Exclude known footers
    if "Daggerheart SRD" in line:
        return False
    # Exclude lines that are just numbers or page numbers
    if re.match(r"^(\d+|Page \d+)$", line.strip()):
        return False
    # Only consider as heading if not too short and not all numbers
    if len(line.strip()) < 4:
        return False
    # ALL CAPS, but not if it's a footer
    if re.match(r"^[A-Z][A-Z\s\-]+$", line.strip()):
        return True
    # Title Case, but not if it's a footer
    if re.match(r"^[A-Z][a-z]+(\s+[A-Z][a-z]+)*[:?]?$", line.strip()):
        return True
    # Ends with colon or question mark, but not if it's a list item
    if line.strip().endswith(":") or line.strip().endswith("?"):
        # Don't split on lines that are part of a bulleted or numbered list
        if not re.match(r"^(\*|\-|\d+\.)", line.strip()):
            return True
    return False


def extract_page_text(pdf_path: str, page_number: int) -> str:
    with open(pdf_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)
        if 0 <= page_number < len(reader.pages):
            return reader.pages[page_number].extract_text() or ""
        else:
            return ""

# ========== TOC PAGE FINDER ==========
def find_toc_page(pdf_path, max_search_pages=10):
    toc_keywords = ["table of contents", "contents", "index", "summary"]
    with open(pdf_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)
        num_pages = min(len(reader.pages), max_search_pages)
        for i in range(num_pages):
            page_text = reader.pages[i].extract_text()
            if page_text:
                page_text_lower = page_text.lower()
                if any(keyword in page_text_lower for keyword in toc_keywords):
                    return i
    return None

def extract_section_headings(toc_text):
    headings = []
    for line in toc_text.split('\n'):
        line = line.strip()
        if not line or line.isdigit() or line.upper() == "CONTENTS":
            continue
        match = re.match(r"^(.*?)(?:\s*[\.\s]+)?(\d+)$", line)
        if match:
            heading = match.group(1).strip()
            if heading:
                headings.append(heading)
    return headings

# ========== PDF SECTION CHUNKING WITH PAGE TAGS ==========
def extract_sections_from_pdf(pdf_path: str, headings: list):
    with open(pdf_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)
        heading_pages = {}
        normalized_headings = [h.strip().lower() for h in headings]
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            lines = text.split('\n')
            for line in lines:
                norm_line = line.strip().lower()
                if norm_line in normalized_headings:
                    idx = normalized_headings.index(norm_line)
                    heading_pages[headings[idx]] = page_num + 1

        full_text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                full_text += page_text + "\n"

        sections = []
        normalized_headings = [h.strip().lower() for h in headings]
        lines = full_text.split('\n')
        current_section = []
        current_title = None
        for line in lines:
            norm_line = line.strip().lower()
            if norm_line in normalized_headings:
                if current_section and current_title:
                    page_num = heading_pages.get(current_title, None)
                    sections.append(({"Section": current_title, "Page": page_num}, "\n".join(current_section).strip()))
                current_title = line.strip()
                current_section = []
            else:
                current_section.append(line)
        if current_section and current_title:
            page_num = heading_pages.get(current_title, None)
            sections.append(({"Section": current_title, "Page": page_num}, "\n".join(current_section).strip()))
        return sections

# ========== STAT BLOCK CHUNKING FOR ADVERSARIES.PDF ==========
def chunk_stat_blocks_from_pdf(pdf_path: str):
    logger.debug('Stat Block Started')
    logger.debug(f"Chunking stat blocks from PDF: {pdf_path}")
    stat_blocks = []
    name_pattern = re.compile(r'^[A-Z][A-Z\s\-]+[A-Z]$')
    skip_words = {"FEATURES"}

    with open(pdf_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            lines = text.split('\n')
            current_block = []
            current_name = None
            for line in lines:
                line_stripped = line.strip()
                if name_pattern.match(line_stripped) and line_stripped not in skip_words:
                    if current_block and current_name:
                        stat_blocks.append({
                            "name": current_name,
                            "page": page_num + 1,
                            "text": "\n".join(current_block).strip()
                        })
                    current_name = line_stripped
                    current_block = [line_stripped]
                else:
                    if current_block is not None:
                        current_block.append(line)
            if current_block and current_name:
                stat_blocks.append({
                    "name": current_name,
                    "page": page_num + 1,
                    "text": "\n".join(current_block).strip()
                })
    return stat_blocks

# ========== DOMAIN CARD CHUNKING FROM PDF ==========
def chunk_domain_cards_from_pdf(pdf_path: str):
    logger.debug('Domain Card Started')
    logger.debug(f"Chunking domain cards from PDF: {pdf_path}")
    card_blocks = []
    domain_pattern = re.compile(r'^[A-Z\s]+DOMAIN$')
    card_pattern = re.compile(r'^(■\s*)?([A-Z][A-Z\s\-]+[A-Z])$')
    skip_words = {"APPENDIX", "DOMAIN CARD REFERENCE"}

    with open(pdf_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            lines = text.split('\n')
            current_domain = None
            current_card = None
            current_block = []
            for line in lines:
                line_stripped = line.strip()
                # Detect domain
                if domain_pattern.match(line_stripped):
                    current_domain = line_stripped
                    continue
                # Detect card/spell name
                if card_pattern.match(line_stripped) and line_stripped not in skip_words:
                    # Save previous card
                    if current_block and current_card:
                        card_blocks.append({
                            "domain": current_domain,
                            "name": current_card,
                            "page": page_num + 1,
                            "text": "\n".join(current_block).strip()
                        })
                    current_card = card_pattern.match(line_stripped).group(2)
                    current_block = [line_stripped]
                else:
                    if current_block is not None:
                        current_block.append(line)
            # Save last card on the page
            if current_block and current_card:
                card_blocks.append({
                    "domain": current_domain,
                    "name": current_card,
                    "page": page_num + 1,
                    "text": "\n".join(current_block).strip()
                })
    return card_blocks

# ========== DOCX SECTION CHUNKING ==========
def extract_sections_from_docx(docx_path: str):
    doc = Document(docx_path)
    sections = []
    current_heading = None
    current_section = []
    for para in doc.paragraphs:
        style = para.style.name if para.style else ""
        text = para.text.strip()
        if not text:
            continue
        if style.startswith("Heading"):
            if current_section and current_heading:
                sections.append(({"Section": current_heading}, "\n".join(current_section)))
            current_heading = text
            current_section = []
        else:
            current_section.append(text)
    if current_section and current_heading:
        sections.append(({"Section": current_heading}, "\n".join(current_section)))
    return sections

# ========== TABLE (CSV/XLSX) EXTRACTION ==========
def extract_chunks_from_table(
    file_path: str,
    heading_columns: Optional[list] = None,
    filter_dict: Optional[dict] = None
):
    ext = os.path.splitext(file_path)[1].lower()
    chunks = []

    if ext == ".csv":
        with open(file_path, newline='', encoding='utf-8') as csvfile:
            reader = csv.DictReader(csvfile)
            for row in reader:
                if filter_dict:
                    skip = False
                    for k, v in filter_dict.items():
                        if row.get(k) != v:
                            skip = True
                            break
                    if skip:
                        continue
                headings = {col: row.get(col) for col in heading_columns} if heading_columns else None
                text = ', '.join(f"{k}: {v}" for k, v in row.items())
                chunks.append((headings, text))
    elif ext == ".xlsx":
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(h) for h in rows[0]]
            for row in rows[1:]:
                row_dict = {headers[i]: str(row[i]) if row[i] is not None else "" for i in range(len(headers))}
                if filter_dict:
                    skip = False
                    for k, v in filter_dict.items():
                        if row_dict.get(k) != v:
                            skip = True
                            break
                    if skip:
                        continue
                headings = {col: row_dict.get(col) for col in heading_columns} if heading_columns else None
                text = ', '.join(f"{k}: {v}" for k, v in row_dict.items())
                chunks.append((headings, text))
    else:
        raise ValueError("Unsupported table file type")
    return chunks

# ========== EMBEDDING ==========
def get_embeddings(texts, embedder):
    return embedder.encode(texts, normalize_embeddings=True)

# ========== OUTPUT PATHS ==========
def get_output_paths(source_path, embeddings_dir=EMBEDDINGS_DIR, chunks_dir=CHUNKS_DIR):
    base = os.path.splitext(os.path.basename(source_path))[0]
    embeddings_path = os.path.join(embeddings_dir, f"{base}_embeddings.npy")
    chunks_path = os.path.join(chunks_dir, f"{base}_chunks.json")
    return embeddings_path, chunks_path
def chunk_by_headings_sequential(full_text, headings):
    lines = full_text.split('\n')
    chunks = []
    last_end = 0
    for i, row in enumerate(headings):
        # Build the lowest available heading for this row
        heading = row['subsubsubsubsection'] or row['subsubsubsection'] or row['subsubsection'] or row['subsection'] or row['section']
        if not heading:
            continue
        # Search for the heading after the last chunk's end
        start = next((j for j in range(last_end, len(lines)) if lines[j].strip().lower() == heading.strip().lower()), None)
        if start is None:
            logger.warning(f"Warning: Heading '{heading}' not found after line {last_end}")
            continue
        # End at the next heading (after start)
        next_start = None
        for j in range(i+1, len(headings)):
            next_heading = headings[j]['subsubsubsubsection'] or headings[j]['subsubsubsection'] or headings[j]['subsubsection'] or headings[j]['subsection'] or headings[j]['section']
            if not next_heading:
                continue
            next_start = next((k for k in range(start+1, len(lines)) if lines[k].strip().lower() == next_heading.strip().lower()), None)
            if next_start is not None:
                break
        end = next_start if next_start is not None else len(lines)
        chunk_text = "\n".join(lines[start:end])
        chunk = {
            "index": i,
            "section": row['section'],
            "subsection": row['subsection'],
            "subsubsection": row['subsubsection'],
            "subsubsubsection": row['subsubsubsection'],
            "subsubsubsubsection": row['subsubsubsubsection'],
            "start_page": row.get('start page'),
            "end_page": row.get('end page'),
            "heading": heading,
            "text": chunk_text
        }
        chunks.append(chunk)
        last_end = end
    return chunks


def extract_full_text_preserve_titles(pdf_path):
    all_text = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            # layout=True tries to keep reading order
            page_text = page.extract_text(layout=True)
            if page_text:
                all_text.append(page_text)
    return "\n".join(all_text)

import re

def remove_footers_headers(text):
    lines = text.split('\n')
    clean_lines = []
    for line in lines:
        # Remove lines that are just footers/headers
        if re.match(r".*Daggerheart SRD.*Page\s*\d+", line):
            continue
        if line.strip().startswith("Daggerheart SRD"):
            continue
        clean_lines.append(line)
    return "\n".join(clean_lines)
# ========== MAIN WORKFLOW ==========
def embedding_generator(
    FILE_PATH,
    embedder,
    BATCH_SIZE=20,
    headings_csv_path="./source/Daggerheart_context_extended.csv",
    heading_columns=None,
    filter_dict=None
):

    if not os.path.exists(headings_csv_path):
        raise FileNotFoundError(f"CSV not found: {headings_csv_path}")
    else:
        logger.debug(f"Using headings from: {headings_csv_path}")
    EMBEDDINGS_PATH, CHUNKS_PATH = get_output_paths(FILE_PATH)
    os.makedirs(os.path.dirname(EMBEDDINGS_PATH), exist_ok=True)
    os.makedirs(os.path.dirname(CHUNKS_PATH), exist_ok=True)

    ext = os.path.splitext(FILE_PATH)[1].lower()
    base = os.path.splitext(os.path.basename(FILE_PATH))[0]
    logger.debug(f"Processing file: {FILE_PATH}")

    if base == "core_rules":
        headings = read_headings_csv(headings_csv_path)
        logger.debug("Reading plain text file...")
        with open(FILE_PATH, "r", encoding="latin-1") as f:
            full_text = f.read()
        
        logger.debug("Splitting by headings...")
        chunk_dicts = chunk_by_headings_sequential(full_text, headings)
        texts = [c["text"] for c in chunk_dicts]
    elif base.lower() == "adversaries" or base.lower() == "environments":
        logger.debug("Special stat block chunking for adversaries.pdf ...")
        stat_blocks = chunk_stat_blocks_from_pdf(FILE_PATH)
        chunk_dicts = []
        for idx, block in enumerate(stat_blocks):
            logger.debug('Processing block:', idx, block["name"])
            chunk_dicts.append({
                "index": idx,
                "name": block["name"],
                "page": block["page"],
                "text": preprocess_text(block["text"])
            })
        texts = [c["text"] for c in chunk_dicts]
    elif base.lower() == "domain_card_reference":
        logger.debug("Special domain card chunking for domain_card_reference.pdf ...")
        card_blocks = chunk_domain_cards_from_pdf(FILE_PATH)
        chunk_dicts = []
        for idx, block in enumerate(card_blocks):
            chunk_dicts.append({
                "index": idx,
                "domain": block["domain"],
                "name": block["name"],
                "page": block["page"],
                "text": preprocess_text(block["text"])
            })
        texts = [c["text"] for c in chunk_dicts]
    elif ext in [".csv", ".xlsx"]:
        logger.debug("Extracting and chunking table (CSV/XLSX)...")
        raw_sections = extract_chunks_from_table(
            FILE_PATH,
            heading_columns=heading_columns,
            filter_dict=filter_dict
        )
        raw_sections = [(h, preprocess_text(s)) for h, s in raw_sections]
        chunk_dicts = []
        for idx, (headings, text) in enumerate(raw_sections):
            chunk_dicts.append({
                "index": idx,
                "headings": headings,
                "text": text
            })
        texts = [c["text"] for c in chunk_dicts]
    elif ext == ".pdf":
        logger.debug("Searching for Table of Contents page...")
        toc_page = None
        if toc_page is None:
            toc_page = find_toc_page(FILE_PATH, max_search_pages=10)
        if toc_page is not None:
            logger.debug(f"Found ToC on page {toc_page+1}")
            toc_text = extract_page_text(FILE_PATH, toc_page)
        else:
            logger.debug("ToC not found, defaulting to page 2.")
            toc_text = extract_page_text(FILE_PATH, 1)
        headings = extract_section_headings(toc_text)
        logger.debug(f"Section headings found in ToC: {headings}")
        section_chunks = extract_sections_from_pdf(FILE_PATH, headings)
        raw_sections = [(h, preprocess_text(s)) for h, s in section_chunks]
        chunk_dicts = []
        for idx, (headings, text) in enumerate(raw_sections):
            chunk_dicts.append({
                "index": idx,
                "headings": headings,
                "text": text
            })
        texts = [c["text"] for c in chunk_dicts]
    elif ext == ".docx":
        logger.debug("Extracting and chunking DOCX by heading...")
        section_chunks = extract_sections_from_docx(FILE_PATH)
        raw_sections = [(h, preprocess_text(s)) for h, s in section_chunks]
        chunk_dicts = []
        for idx, (headings, text) in enumerate(raw_sections):
            chunk_dicts.append({
                "index": idx,
                "headings": headings,
                "text": text
            })
        texts = [c["text"] for c in chunk_dicts]
    elif ext == ".txt":
        logger.debug("Extracting and chunking TXT...")
        full_text = extract_text_from_txt(FILE_PATH)
        chunk_dicts = [{
            "index": 0,
            "headings": None,
            "text": preprocess_text(full_text)
        }]
        texts = [chunk_dicts[0]["text"]]
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    logger.debug(f"Chunk count: {len(chunk_dicts)}")
    for i, chunk in enumerate(chunk_dicts):
        logger.debug(
            f"Chunk {i}: {len(chunk['text']) // 4} tokens (approx), "
            f"Meta: {chunk.get('section')}, "
            f"Subsections: {chunk.get('subsection')}, {chunk.get('subsubsection')}, {chunk.get('subsubsubsection')}, {chunk.get('subsubsubsubsection')}"
        )

    # Generate embeddings in batches
    logger.debug("Generating embeddings...")
    embeddings = []
    for i in range(0, len(chunk_dicts), BATCH_SIZE):
        batch_texts = texts[i:i+BATCH_SIZE]
        batch_embeddings = get_embeddings(batch_texts, embedder)
        embeddings.append(batch_embeddings)
        logger.debug(f"Embedded batch {i//BATCH_SIZE + 1}/{(len(chunk_dicts)-1)//BATCH_SIZE + 1}")
    embeddings = np.vstack(embeddings)

    # Save to disk
    logger.debug(f"Saving embeddings to {EMBEDDINGS_PATH} ...")
    np.save(EMBEDDINGS_PATH, embeddings)
    logger.debug(f"Saving chunks to {CHUNKS_PATH} ...")
    with open(CHUNKS_PATH, "w", encoding="utf-8") as f:
        json.dump(chunk_dicts, f, ensure_ascii=False, indent=2)

    logger.info("Done! You can now load these files in your search app.")

if __name__ == "__main__":
    start = time.perf_counter()
    model = EMBEDDING_MODEL
    embedder = SentenceTransformer(model)

    # Example: process all files in ./source
    source_dir = SOURCE_DIR
    skip_files = {
        'DH-SRD-1.0-June-26-2025.pdf',
        'Daggerheart_context_extended.csv',
        'core_rules.docx',
        'core_rules.pdf',
    }  # Add any files you want to skip
    headings_csv_path = "./source/Daggerheart_context_extended.csv"
    if not os.path.exists(headings_csv_path):
        raise FileNotFoundError(f"CSV not found: {headings_csv_path}")
    else:
        logger.debug(f"Using headings from: {headings_csv_path}")
    for fname in os.listdir(source_dir):
        if not fname.lower().endswith((".pdf", ".docx", ".txt", ".csv", ".xlsx")):
            continue
        if fname in skip_files:
            logger.debug(f"Skipping {fname}")
            continue
        FILE_PATH = os.path.join(source_dir, fname)
        embedding_generator(
            FILE_PATH,
            embedder=embedder,
            headings_csv_path="./source/Daggerheart_context_extended.csv"
        )
    end = time.perf_counter()
    logger.debug(f"Script took {end - start:.2f} seconds")